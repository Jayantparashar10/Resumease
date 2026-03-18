import os
import tempfile
import uuid
from datetime import datetime

import httpx
from dotenv import load_dotenv
from docx import Document


load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env"))

BASE_URL = os.getenv("BASE_URL", "http://127.0.0.1:8000")
SUPABASE_URL = os.getenv("SUPABASE_URL", "").rstrip("/")
SUPABASE_ANON_KEY = os.getenv("SUPABASE_ANON_KEY", "")
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "")
GOOGLE_ID_TOKEN = os.getenv("GOOGLE_ID_TOKEN", "")


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def _log(step: str, status: str, detail: str = "") -> None:
    suffix = f" | {detail}" if detail else ""
    print(f"[{status}] {step}{suffix}")


def _supabase_headers(service: bool = False) -> dict:
    key = SUPABASE_SERVICE_ROLE_KEY if service else SUPABASE_ANON_KEY
    return {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
    }


async def _create_supabase_user_token() -> tuple[str, str]:
    email = f"resumease_{uuid.uuid4().hex[:10]}@resumase.dev"
    password = f"TmpPass_{uuid.uuid4().hex[:12]}"

    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.post(
            f"{SUPABASE_URL}/auth/v1/signup",
            headers=_supabase_headers(service=False),
            json={"email": email, "password": password},
        )
    _assert(r.status_code < 400, f"Supabase signup failed: {r.text}")
    data = r.json()
    token = data.get("access_token")
    user = data.get("user") or {}
    _assert(bool(token), "Supabase signup did not return access_token")
    _assert(bool(user.get("id")), "Supabase signup did not return user id")
    return token, user["id"]


async def _promote_profile_role(user_id: str, role: str, onboarding_data: dict) -> None:
    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.patch(
            f"{SUPABASE_URL}/rest/v1/profiles",
            params={"user_id": f"eq.{user_id}", "select": "user_id"},
            headers={**_supabase_headers(service=True), "Prefer": "return=representation"},
            json={
                "role": role,
                "onboarding_completed": True,
                "onboarding_data": onboarding_data,
                "updated_at": datetime.utcnow().isoformat(),
            },
        )
    _assert(r.status_code < 400, f"Profile role patch failed: {r.text}")


async def main() -> None:
    _assert(bool(SUPABASE_URL), "SUPABASE_URL missing in backend/.env")
    _assert(bool(SUPABASE_ANON_KEY), "SUPABASE_ANON_KEY missing in backend/.env")
    _assert(bool(SUPABASE_SERVICE_ROLE_KEY), "SUPABASE_SERVICE_ROLE_KEY missing in backend/.env")

    async with httpx.AsyncClient(timeout=60) as client:
        # Basic health/docs checks
        r = await client.get(f"{BASE_URL}/health")
        _assert(r.status_code == 200, f"/health failed: {r.text}")
        _log("GET /health", "PASS")

        r = await client.get(f"{BASE_URL}/docs")
        _assert(r.status_code == 200, f"/docs failed: {r.status_code}")
        _log("GET /docs", "PASS")

        access_token = None
        user_id = None

        # Test auth/google endpoint
        if GOOGLE_ID_TOKEN:
            r = await client.post(
                f"{BASE_URL}/api/v1/auth/google",
                json={"id_token": GOOGLE_ID_TOKEN},
            )
            _assert(r.status_code == 200, f"POST /auth/google failed: {r.text}")
            data = r.json()
            access_token = data["access_token"]
            user_id = data["user"]["id"]
            _log("POST /api/v1/auth/google", "PASS", "real Google token")
        else:
            r = await client.post(
                f"{BASE_URL}/api/v1/auth/google",
                json={"id_token": "invalid-token"},
            )
            _assert(r.status_code == 401, "POST /auth/google should reject invalid token")
            _log("POST /api/v1/auth/google", "PASS", "invalid token rejected (set GOOGLE_ID_TOKEN for success-path)")

            access_token, user_id = await _create_supabase_user_token()
            _log("Supabase signup fallback", "PASS", "token acquired for protected endpoint tests")

        auth_headers = {"Authorization": f"Bearer {access_token}"}

        # Onboarding (student)
        student_payload = {
            "role": "student",
            "data": {
                "college": "Test Institute",
                "degree": "B.Tech",
                "graduation_year": 2026,
                "target_roles": ["Backend Developer"],
                "skills_self_reported": ["Python", "FastAPI"],
            },
        }
        r = await client.post(
            f"{BASE_URL}/api/v1/auth/onboarding",
            headers=auth_headers,
            json=student_payload,
        )
        _assert(r.status_code == 200, f"POST /auth/onboarding failed: {r.text}")
        _log("POST /api/v1/auth/onboarding", "PASS")

        # Promote to recruiter to create jobs
        recruiter_onboarding = {
            "company": "Acme Inc",
            "designation": "HR",
            "hiring_for": ["Backend Developer"],
            "company_size": "51-200",
        }
        await _promote_profile_role(user_id, "recruiter", recruiter_onboarding)
        _log("Profile switch -> recruiter", "PASS")

        # Recruiter job create
        job_payload = {
            "title": "Backend Engineer",
            "company": "Acme Inc",
            "description": "Need FastAPI and SQL experience",
            "required_skills": ["Python", "FastAPI", "SQL"],
            "location": "Remote",
            "experience_years": 1,
        }
        r = await client.post(
            f"{BASE_URL}/api/v1/jobs/create",
            headers=auth_headers,
            json=job_payload,
        )
        _assert(r.status_code == 201, f"POST /jobs/create failed: {r.text}")
        created_job = r.json()
        job_id = created_job["id"]
        _log("POST /api/v1/jobs/create", "PASS", f"job_id={job_id}")

        # Recruiter job list/get/update
        r = await client.get(f"{BASE_URL}/api/v1/jobs/list", headers=auth_headers)
        _assert(r.status_code == 200, f"GET /jobs/list failed: {r.text}")
        _log("GET /api/v1/jobs/list", "PASS")

        r = await client.get(f"{BASE_URL}/api/v1/jobs/{job_id}", headers=auth_headers)
        _assert(r.status_code == 200, f"GET /jobs/{{id}} failed: {r.text}")
        _log("GET /api/v1/jobs/{job_id}", "PASS")

        updated_payload = {
            **job_payload,
            "title": "Backend Engineer II",
        }
        r = await client.put(
            f"{BASE_URL}/api/v1/jobs/{job_id}",
            headers=auth_headers,
            json=updated_payload,
        )
        _assert(r.status_code == 200, f"PUT /jobs/{{id}} failed: {r.text}")
        _log("PUT /api/v1/jobs/{job_id}", "PASS")

        # Switch back to student for resume and ATS flow
        await _promote_profile_role(user_id, "student", student_payload["data"])
        _log("Profile switch -> student", "PASS")

        # Create sample DOCX for upload
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp_path = tmp.name
        doc = Document()
        doc.add_heading("Resume", 0)
        doc.add_paragraph("John Doe")
        doc.add_paragraph("Email: john@example.com")
        doc.add_paragraph("Skills: Python, FastAPI, SQL")
        doc.add_paragraph("GitHub: https://github.com/johndoe")
        doc.save(tmp_path)

        with open(tmp_path, "rb") as f:
            files = {
                "file": (
                    "resume.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            }
            r = await client.post(
                f"{BASE_URL}/api/v1/resumes/upload",
                headers=auth_headers,
                files=files,
            )
        _assert(r.status_code == 201, f"POST /resumes/upload failed: {r.text}")
        resume = r.json()
        resume_id = resume["id"]
        _log("POST /api/v1/resumes/upload", "PASS", f"resume_id={resume_id}")

        # Student resume list/get
        r = await client.get(f"{BASE_URL}/api/v1/resumes/list", headers=auth_headers)
        _assert(r.status_code == 200, f"GET /resumes/list failed: {r.text}")
        _log("GET /api/v1/resumes/list", "PASS")

        r = await client.get(f"{BASE_URL}/api/v1/resumes/{resume_id}", headers=auth_headers)
        _assert(r.status_code == 200, f"GET /resumes/{{id}} failed: {r.text}")
        _log("GET /api/v1/resumes/{resume_id}", "PASS")

        # ATS score + history
        r = await client.post(
            f"{BASE_URL}/api/v1/ats/score",
            headers=auth_headers,
            json={"resume_id": resume_id, "job_id": job_id},
        )
        _assert(r.status_code == 201, f"POST /ats/score failed: {r.text}")
        score = r.json()
        score_id = score["id"]
        _log("POST /api/v1/ats/score", "PASS", f"score_id={score_id}")

        r = await client.get(f"{BASE_URL}/api/v1/ats/score/{score_id}", headers=auth_headers)
        _assert(r.status_code == 200, f"GET /ats/score/{{id}} failed: {r.text}")
        _log("GET /api/v1/ats/score/{score_id}", "PASS")

        r = await client.get(f"{BASE_URL}/api/v1/ats/history", headers=auth_headers)
        _assert(r.status_code == 200, f"GET /ats/history failed: {r.text}")
        _log("GET /api/v1/ats/history", "PASS")

        # Switch to recruiter and test candidates + delete job
        await _promote_profile_role(user_id, "recruiter", recruiter_onboarding)
        _log("Profile switch -> recruiter", "PASS")

        r = await client.get(
            f"{BASE_URL}/api/v1/recruiter/candidates/{job_id}",
            headers=auth_headers,
        )
        _assert(r.status_code == 200, f"GET /recruiter/candidates/{{job_id}} failed: {r.text}")
        _log("GET /api/v1/recruiter/candidates/{job_id}", "PASS")

        r = await client.delete(f"{BASE_URL}/api/v1/jobs/{job_id}", headers=auth_headers)
        _assert(r.status_code == 204, f"DELETE /jobs/{{id}} failed: {r.text}")
        _log("DELETE /api/v1/jobs/{job_id}", "PASS")

        # Cleanup resume
        await _promote_profile_role(user_id, "student", student_payload["data"])
        r = await client.delete(f"{BASE_URL}/api/v1/resumes/{resume_id}", headers=auth_headers)
        _assert(r.status_code == 204, f"DELETE /resumes/{{id}} failed: {r.text}")
        _log("DELETE /api/v1/resumes/{resume_id}", "PASS")

    print("\nAll smoke tests passed.")


if __name__ == "__main__":
    import asyncio

    asyncio.run(main())
